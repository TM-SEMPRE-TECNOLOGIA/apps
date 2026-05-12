"""
test_word_utils.py - Testes unitários para word_utils.py

Testes para:
- extract_placeholders()
- substitute_placeholders()
- get_template_metadata()
- validate_substitutions()
"""

import pytest
import tempfile
from pathlib import Path
from docx import Document
from word_utils import (
    extract_placeholders,
    substitute_placeholders,
    get_template_metadata,
    validate_substitutions,
    EXPECTED_PLACEHOLDERS
)


# ============================================================================
# FIXTURES
# ============================================================================

@pytest.fixture
def template_path():
    """Retorna caminho do primeiro template disponível"""
    templates_dir = Path(__file__).parent / "templates"
    templates = list(templates_dir.glob("*.docx"))
    if not templates:
        pytest.skip("Nenhum template encontrado")
    return str(templates[0])


@pytest.fixture
def template_name():
    """Retorna nome de um template para teste"""
    return "MODELO - 3575 - TANGARA DA SERRA.docx"


@pytest.fixture
def valid_substitutions():
    """Dados válidos para substituição"""
    return {
        'nr_os': '1753',
        'data_elaboracao': '2026-05-03',
        'data_atendimento': '2026-05-01',
        'agencia_codigo': '3575',
        'agencia_nome': 'TANGARA DA SERRA',
        'endereco': 'Avenida Brasil, 1000 - Tangará da Serra, MT',
        'responsavel_dependencia': '123456 - João Silva'
    }


# ============================================================================
# TESTES: extract_placeholders()
# ============================================================================

class TestExtractPlaceholders:
    """Testes para função extract_placeholders()"""

    def test_extract_placeholders_returns_dict(self, template_path):
        """Deve retornar um dicionário"""
        result = extract_placeholders(template_path)
        assert isinstance(result, dict)

    def test_extract_placeholders_has_required_keys(self, template_path):
        """Deve ter chaves obrigatórias"""
        result = extract_placeholders(template_path)
        
        required_keys = ['placeholders', 'locations', 'missing', 'extra', 'total_found', 'total_expected']
        for key in required_keys:
            assert key in result, f"Chave '{key}' não encontrada"

    def test_extract_placeholders_found_count_valid(self, template_path):
        """total_found deve ser válido"""
        result = extract_placeholders(template_path)
        assert isinstance(result['total_found'], int)
        assert result['total_found'] >= 0
        assert result['total_found'] <= len(EXPECTED_PLACEHOLDERS)

    def test_extract_placeholders_expected_count_matches(self, template_path):
        """total_expected deve ser 7"""
        result = extract_placeholders(template_path)
        assert result['total_expected'] == 7

    def test_extract_placeholders_placeholders_is_set(self, template_path):
        """'placeholders' deve ser um set"""
        result = extract_placeholders(template_path)
        assert isinstance(result['placeholders'], set)

    def test_extract_placeholders_missing_is_set(self, template_path):
        """'missing' deve ser um set"""
        result = extract_placeholders(template_path)
        assert isinstance(result['missing'], set)

    def test_extract_placeholders_extra_is_set(self, template_path):
        """'extra' deve ser um set"""
        result = extract_placeholders(template_path)
        assert isinstance(result['extra'], set)

    def test_extract_placeholders_locations_has_all_found(self, template_path):
        """locations deve ter todas os placeholders encontrados"""
        result = extract_placeholders(template_path)
        
        for placeholder in result['placeholders']:
            assert placeholder in result['locations']
            assert isinstance(result['locations'][placeholder], list)
            assert len(result['locations'][placeholder]) > 0

    def test_extract_placeholders_missing_not_in_placeholders(self, template_path):
        """missing não deve estar em placeholders"""
        result = extract_placeholders(template_path)
        assert result['missing'].isdisjoint(result['placeholders'])

    def test_extract_placeholders_file_not_found(self):
        """Deve lançar exceção se arquivo não existir"""
        with pytest.raises(FileNotFoundError):
            extract_placeholders("/caminho/inexistente/template.docx")


# ============================================================================
# TESTES: substitute_placeholders()
# ============================================================================

class TestSubstitutePlaceholders:
    """Testes para função substitute_placeholders()"""

    def test_substitute_placeholders_returns_dict(self, template_path, valid_substitutions):
        """Deve retornar um dicionário"""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "output.docx"
            result = substitute_placeholders(
                template_path=template_path,
                output_path=str(output_path),
                substitutions=valid_substitutions
            )
            assert isinstance(result, dict)

    def test_substitute_placeholders_success_true(self, template_path, valid_substitutions):
        """'success' deve ser True para dados válidos"""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "output.docx"
            result = substitute_placeholders(
                template_path=template_path,
                output_path=str(output_path),
                substitutions=valid_substitutions
            )
            assert result['success'] is True

    def test_substitute_placeholders_file_created(self, template_path, valid_substitutions):
        """Arquivo de saída deve ser criado"""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "output.docx"
            substitute_placeholders(
                template_path=template_path,
                output_path=str(output_path),
                substitutions=valid_substitutions
            )
            assert output_path.exists()

    def test_substitute_placeholders_substitutions_made(self, template_path, valid_substitutions):
        """Deve indicar número de substituições realizadas"""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "output.docx"
            result = substitute_placeholders(
                template_path=template_path,
                output_path=str(output_path),
                substitutions=valid_substitutions
            )
            assert result['substitutions_made'] >= 0

    def test_substitute_placeholders_output_path_in_result(self, template_path, valid_substitutions):
        """Resultado deve incluir output_path"""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "output.docx"
            result = substitute_placeholders(
                template_path=template_path,
                output_path=str(output_path),
                substitutions=valid_substitutions
            )
            assert 'output_path' in result
            assert result['output_path'] == str(output_path)

    def test_substitute_placeholders_template_not_found(self, valid_substitutions):
        """Deve lançar exceção se template não existir"""
        with pytest.raises(FileNotFoundError):
            substitute_placeholders(
                template_path="/caminho/inexistente/template.docx",
                output_path="/tmp/output.docx",
                substitutions=valid_substitutions
            )

    def test_substitute_placeholders_creates_parent_directory(self, template_path, valid_substitutions):
        """Deve criar diretório pai se não existir"""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "subdir" / "output.docx"
            substitute_placeholders(
                template_path=template_path,
                output_path=str(output_path),
                substitutions=valid_substitutions
            )
            assert output_path.parent.exists()

    def test_substitute_placeholders_generated_file_is_valid_docx(self, template_path, valid_substitutions):
        """Arquivo gerado deve ser um .docx válido"""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "output.docx"
            substitute_placeholders(
                template_path=template_path,
                output_path=str(output_path),
                substitutions=valid_substitutions
            )
            # Tentar abrir com python-docx
            doc = Document(str(output_path))
            assert len(doc.paragraphs) > 0


# ============================================================================
# TESTES: get_template_metadata()
# ============================================================================

class TestGetTemplateMetadata:
    """Testes para função get_template_metadata()"""

    def test_get_template_metadata_returns_dict(self, template_name):
        """Deve retornar um dicionário"""
        result = get_template_metadata(template_name)
        assert isinstance(result, dict)

    def test_get_template_metadata_has_required_keys(self, template_name):
        """Deve ter chaves obrigatórias"""
        result = get_template_metadata(template_name)
        
        required_keys = ['name', 'code', 'agency', 'placeholders', 'fields', 
                        'dynamic_fields', 'fixed_fields']
        for key in required_keys:
            assert key in result, f"Chave '{key}' não encontrada"

    def test_get_template_metadata_extracts_code_from_name(self):
        """Deve extrair código do nome do template"""
        result = get_template_metadata("MODELO - 3575 - TANGARA DA SERRA.docx")
        assert result['code'] == '3575'

    def test_get_template_metadata_extracts_agency_from_name(self):
        """Deve extrair nome da agência do template"""
        result = get_template_metadata("MODELO - 3575 - TANGARA DA SERRA.docx")
        assert result['agency'] == 'TANGARA DA SERRA'

    def test_get_template_metadata_placeholders_has_7_items(self, template_name):
        """placeholders deve ter 7 itens"""
        result = get_template_metadata(template_name)
        assert len(result['placeholders']) == 7

    def test_get_template_metadata_dynamic_fields_is_list(self, template_name):
        """dynamic_fields deve ser uma lista"""
        result = get_template_metadata(template_name)
        assert isinstance(result['dynamic_fields'], list)
        assert len(result['dynamic_fields']) == 3

    def test_get_template_metadata_fixed_fields_is_list(self, template_name):
        """fixed_fields deve ser uma lista"""
        result = get_template_metadata(template_name)
        assert isinstance(result['fixed_fields'], list)
        assert len(result['fixed_fields']) == 4

    def test_get_template_metadata_fields_describes_all_placeholders(self, template_name):
        """fields deve descrever todos os placeholders"""
        result = get_template_metadata(template_name)
        for placeholder in result['placeholders']:
            assert placeholder in result['fields']

    def test_get_template_metadata_field_has_label(self, template_name):
        """Cada campo deve ter label"""
        result = get_template_metadata(template_name)
        for field_name, field_config in result['fields'].items():
            assert 'label' in field_config
            assert isinstance(field_config['label'], str)

    def test_get_template_metadata_field_has_type(self, template_name):
        """Cada campo deve ter type (string ou date)"""
        result = get_template_metadata(template_name)
        for field_name, field_config in result['fields'].items():
            assert 'type' in field_config
            assert field_config['type'] in ['string', 'date']


# ============================================================================
# TESTES: validate_substitutions()
# ============================================================================

class TestValidateSubstitutions:
    """Testes para função validate_substitutions()"""

    def test_validate_substitutions_returns_dict(self, template_name, valid_substitutions):
        """Deve retornar um dicionário"""
        result = validate_substitutions(template_name, valid_substitutions)
        assert isinstance(result, dict)

    def test_validate_substitutions_valid_data_returns_true(self, template_name, valid_substitutions):
        """Dados válidos devem retornar valid=True"""
        result = validate_substitutions(template_name, valid_substitutions)
        assert result['valid'] is True

    def test_validate_substitutions_valid_data_no_errors(self, template_name, valid_substitutions):
        """Dados válidos não devem ter erros"""
        result = validate_substitutions(template_name, valid_substitutions)
        assert len(result['errors']) == 0

    def test_validate_substitutions_missing_required_field(self, template_name, valid_substitutions):
        """Falta de campo obrigatório deve ser erro"""
        invalid_data = valid_substitutions.copy()
        del invalid_data['nr_os']
        
        result = validate_substitutions(template_name, invalid_data)
        assert result['valid'] is False
        assert len(result['errors']) > 0

    def test_validate_substitutions_empty_field_is_error(self, template_name, valid_substitutions):
        """Campo vazio deve ser erro"""
        invalid_data = valid_substitutions.copy()
        invalid_data['nr_os'] = ''
        
        result = validate_substitutions(template_name, invalid_data)
        assert result['valid'] is False

    def test_validate_substitutions_invalid_date_format(self, template_name, valid_substitutions):
        """Data em formato inválido deve ser erro"""
        invalid_data = valid_substitutions.copy()
        invalid_data['data_atendimento'] = 'abc'
        
        result = validate_substitutions(template_name, invalid_data)
        assert result['valid'] is False
        assert len(result['errors']) > 0

    def test_validate_substitutions_has_required_result_keys(self, template_name, valid_substitutions):
        """Resultado deve ter chaves obrigatórias"""
        result = validate_substitutions(template_name, valid_substitutions)
        
        required_keys = ['valid', 'errors', 'warnings', 'missing', 
                        'fields_provided', 'fields_expected']
        for key in required_keys:
            assert key in result, f"Chave '{key}' não encontrada"

    def test_validate_substitutions_errors_is_list(self, template_name, valid_substitutions):
        """errors deve ser uma lista"""
        result = validate_substitutions(template_name, valid_substitutions)
        assert isinstance(result['errors'], list)

    def test_validate_substitutions_warnings_is_list(self, template_name, valid_substitutions):
        """warnings deve ser uma lista"""
        result = validate_substitutions(template_name, valid_substitutions)
        assert isinstance(result['warnings'], list)


# ============================================================================
# TESTES DE INTEGRAÇÃO
# ============================================================================

class TestIntegration:
    """Testes de integração entre funções"""

    def test_extract_and_validate_consistency(self, template_path, template_name, valid_substitutions):
        """extract_placeholders e validate_substitutions devem ser consistentes"""
        extracted = extract_placeholders(template_path)
        validated = validate_substitutions(template_name, valid_substitutions)
        
        # Se extraído encontrou todos os placeholders esperados,
        # validação com dados completos também deve passar
        if extracted['total_found'] == extracted['total_expected']:
            assert validated['valid'] is True

    def test_full_workflow(self, template_path, template_name, valid_substitutions):
        """Workflow completo: extract → validate → substitute"""
        # 1. Extract
        extracted = extract_placeholders(template_path)
        assert extracted['total_found'] > 0
        
        # 2. Validate
        validated = validate_substitutions(template_name, valid_substitutions)
        if extracted['total_found'] == extracted['total_expected']:
            assert validated['valid'] is True
        
        # 3. Substitute
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "output.docx"
            result = substitute_placeholders(
                template_path=template_path,
                output_path=str(output_path),
                substitutions=valid_substitutions
            )
            assert result['success'] is True
            assert output_path.exists()


# ============================================================================
# MAIN
# ============================================================================

if __name__ == '__main__':
    pytest.main([__file__, '-v', '--tb=short'])
